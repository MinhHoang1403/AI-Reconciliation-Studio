"""F01–F04: inspect files, preserve cell values and attach exact source positions.

CSV is read without inference by Polars. XLSX uses fastexcel for structural inspection,
then XML and openpyxl metadata to preserve numeric lexical values, styles and formulas.
Header selection is a conservative [D] heuristic; ambiguous regions require explicit review.
"""

import csv
import hashlib
import io
import posixpath
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import BinaryIO, Literal

import fastexcel
import openpyxl
import polars as pl
from defusedxml import ElementTree

from studio.config import Settings
from studio.contracts import Contract, SourceColumn, WorkspacePolicy
from studio.errors import DomainError

PARSER_REVISION = "tabular-1"
MAIN_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
REL_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


class ParseConfig(Contract):
    encoding: Literal["utf-8-sig", "cp1258"] = "utf-8-sig"
    delimiter: Literal[",", ";", "\t", "|"] | None = None
    header_row: int | None = None
    include_hidden: bool = False


class FileInspection(Contract):
    sha256: str
    filename: str
    media_type: str
    size: int
    extension: str


@dataclass
class ParsedFrame:
    frame: pl.DataFrame
    sheet: str
    header_row: int
    columns: list[SourceColumn]
    warnings: list[str]
    cell_metadata: list[dict[str, str | int | None]]
    preamble: list[list[str | None]]


def inspect_upload(
    stream: BinaryIO, filename: str, content_type: str, config: Settings
) -> FileInspection:
    """F01: validate finite seekable upload; return SHA256 without changing raw bytes.

    Reject path names, type spoofing, encrypted/macros/unsafe archives and XML entities.
    Never execute formulas or embedded content. Limits apply before uncompressed parsing.
    """
    if not filename or len(filename) > 255 or re.search(r"[/\\:\x00-\x1f]", filename):
        raise DomainError("INVALID_FILENAME", "Tên tệp không hợp lệ.")
    ext = Path(filename).suffix.lower()
    allowed = {
        ".csv": {"text/csv", "text/plain", "application/vnd.ms-excel"},
        ".xlsx": {"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"},
        ".pdf": {"application/pdf"},
        ".docx": {"application/vnd.openxmlformats-officedocument.wordprocessingml.document"},
    }
    if ext not in allowed or content_type not in allowed[ext] | {"application/octet-stream"}:
        raise DomainError("UNSUPPORTED_FILE_TYPE", "Chỉ hỗ trợ CSV, XLSX, PDF và DOCX hợp lệ.")
    digest = hashlib.sha256()
    size = 0
    stream.seek(0)
    while block := stream.read(1024 * 1024):
        size += len(block)
        if size > config.max_upload_bytes:
            raise DomainError("FILE_TOO_LARGE", "Tệp vượt quá giới hạn dung lượng.", 413)
        digest.update(block)
    if size == 0:
        raise DomainError("EMPTY_FILE", "Tệp không có dữ liệu.")
    stream.seek(0)
    prefix = stream.read(8192)
    stream.seek(0)
    if ext == ".csv" and (b"\x00" in prefix or prefix.startswith((b"PK", b"%PDF"))):
        raise DomainError("FILE_SIGNATURE_MISMATCH", "Nội dung không phải CSV dạng văn bản.")
    if ext == ".pdf" and not prefix.startswith(b"%PDF-"):
        raise DomainError("FILE_SIGNATURE_MISMATCH", "Nội dung không phải PDF hợp lệ.")
    if ext in {".xlsx", ".docx"}:
        try:
            with zipfile.ZipFile(stream) as archive:
                members = archive.infolist()
                names = [m.filename for m in members]
                required = "xl/workbook.xml" if ext == ".xlsx" else "word/document.xml"
                if required not in names or len(members) > 5000 or len(names) != len(set(names)):
                    raise ValueError("invalid package")
                if sum(m.file_size for m in members) > config.max_zip_bytes:
                    raise ValueError("expanded size")
                for member in members:
                    if (
                        member.flag_bits & 1
                        or ".." in member.filename.split("/")
                        or member.filename.startswith("/")
                        or "\\" in member.filename
                        or re.search(
                            r"vbaproject|activex|embeddings|externallinks", member.filename, re.I
                        )
                    ):
                        raise ValueError("unsafe package")
                    if member.filename.endswith((".xml", ".rels")):
                        ElementTree.fromstring(archive.read(member))
                if archive.testzip() is not None:
                    raise ValueError("CRC mismatch")
        except Exception as exc:
            raise DomainError(
                "INVALID_OFFICE_ARCHIVE", "Tệp Office hỏng hoặc chứa nội dung bị chặn."
            ) from exc
        finally:
            stream.seek(0)
    return FileInspection(
        sha256=digest.hexdigest(),
        filename=filename,
        media_type=content_type,
        size=size,
        extension=ext,
    )


def choose_parser(inspection: FileInspection, policy: WorkspacePolicy) -> str:
    """F02: local tabular route first; document cloud must have explicit allowance."""
    if inspection.extension in {".csv", ".xlsx"}:
        return "tabular"
    if not policy.allow_cloud_parsing or policy.parser_request_budget < 1:
        raise DomainError("CLOUD_PARSING_DISABLED", "PDF/DOCX cần quyền và ngân sách parser cloud.")
    return "llamaparse"


def _header_index(rows: list[list[str | None]], requested: int | None) -> tuple[int, list[str]]:
    if requested is not None:
        if requested < 1 or requested > len(rows):
            raise DomainError("INVALID_HEADER_ROW", "Dòng tiêu đề nằm ngoài bảng.")
        return requested - 1, []
    for index, row in enumerate(rows[:30]):
        populated = [v for v in row if v not in (None, "")]
        if len(populated) >= 2 and len(set(populated)) == len(populated):
            return index, ["HEADER_INFERRED_REVIEW"] if index else []
    if rows and any(v not in (None, "") for v in rows[0]):
        return 0, ["SINGLE_COLUMN_HEADER_REVIEW"]
    raise DomainError("HEADER_NOT_FOUND", "Không tìm thấy tiêu đề; hãy chọn dòng tiêu đề.")


def _columns(header: list[str | None]) -> tuple[list[SourceColumn], list[str]]:
    names = [value or f"Cột {i + 1}" for i, value in enumerate(header)]
    warnings = []
    if len(names) != len(set(names)) or any(value in (None, "") for value in header):
        warnings.append("EMPTY_OR_DUPLICATE_HEADER")
    return [
        SourceColumn(id=f"c{i + 1}", name=n, source_index=i + 1) for i, n in enumerate(names)
    ], warnings


def _csv(path: Path, parse_config: ParseConfig, config: Settings) -> list[ParsedFrame]:
    try:
        # Decode strictly. Alternative encoding is an explicit user setting, never silent fallback.
        text = path.read_text(encoding=parse_config.encoding, errors="strict")
        delimiter: str = parse_config.delimiter or ""
        if not delimiter:
            try:
                delimiter = csv.Sniffer().sniff(text[:32768], delimiters=",;\t|").delimiter
            except csv.Error:
                delimiter = ","
        reader = csv.reader(io.StringIO(text), delimiter=delimiter, strict=True)
        rows: list[list[str | None]] = []
        source_rows: list[int] = []
        previous_line = 0
        for record in reader:
            source_rows.append(previous_line + 1)
            rows.append(list(record))
            previous_line = reader.line_num
        if (
            len(rows) > config.max_rows + 30
            or max((len(r) for r in rows), default=0) > config.max_columns
        ):
            raise DomainError("TABLE_LIMIT_EXCEEDED", "Bảng vượt giới hạn số dòng hoặc cột.")
        head, warnings = _header_index(rows, parse_config.header_row)
        columns, more = _columns(rows[head])
        # Polars preserves quoted empty string versus absent/null. No numeric inference.
        frame = pl.read_csv(
            text.encode("utf-8"),
            separator=delimiter,
            skip_rows=head,
            infer_schema=False,
            new_columns=[c.id for c in columns],
            row_index_name="_source_row",
            row_index_offset=head + 2,
            raise_if_empty=True,
        )
        if frame.height != len(rows) - head - 1:
            raise DomainError(
                "CSV_ROW_COUNT_MISMATCH", "Có dòng trống hoặc cấu trúc CSV cần xác nhận."
            )
        frame = frame.with_columns(
            pl.Series("_source_row", source_rows[head + 1 :], dtype=pl.UInt32)
        )
        return [
            ParsedFrame(frame, "CSV", source_rows[head], columns, warnings + more, [], rows[:head])
        ]
    except DomainError:
        raise
    except Exception as exc:
        raise DomainError(
            "CSV_PARSE_ERROR", "Không đọc được CSV; kiểm tra encoding, dấu phân cách và tiêu đề."
        ) from exc


def _xlsx(path: Path, parse_config: ParseConfig, config: Settings) -> list[ParsedFrame]:
    try:
        reader = fastexcel.read_excel(path)
        workbook = openpyxl.load_workbook(path, read_only=False, data_only=False, keep_links=False)
        results = []
        with zipfile.ZipFile(path) as archive:
            book = ElementTree.fromstring(archive.read("xl/workbook.xml"))
            rels = ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
            targets = {r.attrib["Id"]: r.attrib["Target"] for r in rels}
            for sheet_node in book.findall(f"{MAIN_NS}sheets/{MAIN_NS}sheet"):
                name = sheet_node.attrib["name"]
                ws = workbook[name]
                warnings: list[str] = []
                if ws.sheet_state != "visible":
                    warnings.append("HIDDEN_SHEET_REVIEW")
                if ws.max_row > config.max_rows + 30 or ws.max_column > config.max_columns:
                    raise DomainError(
                        "TABLE_LIMIT_EXCEEDED", "Sheet vượt giới hạn số dòng hoặc cột."
                    )
                # Inspect structure with fastexcel; monetary values still use XML lexemes.
                reader.load_sheet(name, header_row=None, n_rows=30, dtypes="string")
                target = targets[sheet_node.attrib[f"{REL_NS}id"]]
                target = (
                    target.lstrip("/")
                    if target.startswith("/")
                    else posixpath.normpath("xl/" + target)
                )
                tree = ElementTree.fromstring(archive.read(target))
                values: dict[str, str | None] = {}
                metadata: list[dict[str, str | int | None]] = []
                max_row, max_col = 0, 0
                for cell in tree.iter(f"{MAIN_NS}c"):
                    address = cell.attrib["r"]
                    pycell = ws[address]
                    max_row = max(max_row, pycell.row)
                    max_col = max(max_col, pycell.column)
                    raw = cell.find(f"{MAIN_NS}v")
                    formula = cell.find(f"{MAIN_NS}f")
                    kind = cell.attrib.get("t", "n")
                    value: str | None = raw.text if raw is not None else None
                    if kind in {"s", "inlineStr", "str"} and formula is None:
                        value = None if pycell.value is None else str(pycell.value)
                    if formula is not None:
                        warnings.append("FORMULA_UNVERIFIED")
                        metadata.append({"cell": address, "formula": formula.text, "cached": value})
                    if kind == "e":
                        warnings.append("CELL_ERROR_UNVERIFIED")
                    # Keep identifier display width as well as the original numeric lexeme.
                    if value is not None and re.fullmatch(r"0{2,}", pycell.number_format):
                        metadata.append(
                            {"cell": address, "numeric_raw": value, "format": pycell.number_format}
                        )
                        if re.fullmatch(r"\d+", value):
                            value = value.zfill(len(pycell.number_format))
                    if pycell.is_date:
                        warnings.append("EXCEL_DATE_SERIAL_REVIEW")
                        metadata.append(
                            {"cell": address, "date_serial": value, "format": pycell.number_format}
                        )
                    values[address] = value
                if max_row == 0:
                    continue
                from openpyxl.utils import get_column_letter

                rows = [
                    [values.get(f"{get_column_letter(c)}{r}") for c in range(1, max_col + 1)]
                    for r in range(1, max_row + 1)
                ]
                head, head_warnings = _header_index(rows, parse_config.header_row)
                columns, column_warnings = _columns(rows[head])
                if ws.merged_cells:
                    warnings.append("MERGED_CELLS_REVIEW")
                    metadata.extend({"merged_range": str(area)} for area in ws.merged_cells.ranges)
                frame = pl.DataFrame(
                    rows[head + 1 :], schema={c.id: pl.String for c in columns}, orient="row"
                )
                frame = frame.with_row_index("_source_row", offset=head + 2)
                results.append(
                    ParsedFrame(
                        frame,
                        name,
                        head + 1,
                        columns,
                        sorted(set(warnings + head_warnings + column_warnings)),
                        metadata,
                        rows[:head],
                    )
                )
        workbook.close()
        if not results:
            raise DomainError("NO_TABLES", "Workbook không có bảng dữ liệu.")
        return results
    except DomainError:
        raise
    except Exception as exc:
        raise DomainError(
            "XLSX_PARSE_ERROR", "Không đọc được XLSX; tệp hoặc vùng bảng không hợp lệ."
        ) from exc


def parse_to_artifacts(
    path: Path, inspection: FileInspection, parse_config: ParseConfig, config: Settings
) -> list[ParsedFrame]:
    """F03: read immutable local source; preserve all post-header rows and raw cell strings."""
    if hashlib.sha256(path.read_bytes()).hexdigest() != inspection.sha256:
        raise DomainError("RAW_CHECKSUM_MISMATCH", "Checksum tệp nguồn đã thay đổi.")
    if inspection.extension == ".csv":
        return _csv(path, parse_config, config)
    if inspection.extension == ".xlsx":
        return _xlsx(path, parse_config, config)
    raise DomainError("PARSER_ROUTE_INVALID", "Định dạng này cần parser tài liệu có cấp phép.")


def verify_parse(parsed: ParsedFrame, config: Settings) -> str:
    """F04: validate structure without repairing values; risky extraction remains unverified."""
    if parsed.frame.height > config.max_rows or not parsed.columns:
        raise DomainError("TABLE_LIMIT_EXCEEDED", "Bảng không đạt giới hạn hoặc thiếu cột.")
    if parsed.frame["_source_row"].n_unique() != parsed.frame.height:
        raise DomainError("SOURCE_ROW_DUPLICATE", "Vị trí dòng nguồn bị trùng.")
    return "UNVERIFIED_EXTRACTION" if parsed.warnings else "STRUCTURE_VERIFIED"
