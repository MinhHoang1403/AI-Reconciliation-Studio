"""CSV/XLSX application export of one frozen MetricSnapshot, with typed cells.

No spreadsheet formulas recalculate business facts. Numeric Decimal cells remain
numeric (including negatives); untrusted strings get text treatment. The OOXML
numeric lexeme is patched from Decimal directly to avoid openpyxl's float formatter.
Excel itself has a 15-digit arithmetic limit, disclosed in the Summary sheet.
"""

import csv
import json
import re
import xml.etree.ElementTree as ET
from decimal import Decimal
from pathlib import Path
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select

from studio.access import Actor
from studio.config import settings
from studio.contracts.reconciliation import MetricSnapshot, ReconciliationContract
from studio.database import ExceptionRow, ExportRow, ReconciliationContractRow, RunRow, identifier
from studio.errors import DomainError
from studio.reconciliation_api import case_view, require_verified
from studio.reconciliation_repository import all_match_groups, load_canonical_records
from studio.storage import checksum, storage_path

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def safe_text(value: str) -> str:
    """CSV text interpretation guard, without coercing legitimate numeric negatives."""
    stripped = re.sub(r"^[\s\x00-\x1f\ufeff]*", "", value)
    return "'" + value if stripped.startswith(("=", "+", "-", "@")) else value


def report_sections(db: Any, run: RunRow) -> dict[str, list[list[Any]]]:
    from studio.contracts import CanonicalDataset

    require_verified(run)
    snapshot = MetricSnapshot.model_validate(run.metric_snapshot)
    contract = db.get(ReconciliationContractRow, run.contract_id)
    assert contract is not None
    rule = ReconciliationContract.model_validate(contract.definition)
    sections: dict[str, list[list[Any]]] = {
        "Summary": [
            ["Field", "Value"],
            ["Run ID", run.id],
            ["Verification", run.status],
            ["Comparison unit", snapshot.comparison_unit],
            ["Eligible definition", snapshot.eligible_definition],
            ["Eligible units", snapshot.eligible_units],
            ["Period start", str(snapshot.period_start)],
            ["Period end exclusive", str(snapshot.period_end)],
            ["Timezone", snapshot.timezone],
            ["Open cases at run", snapshot.unresolved_cases_at_run],
            ["Snapshot JSON", snapshot.model_dump_json()],
            [
                "Precision",
                "Decimal numeric lexemes are exact in XLSX/CSV; "
                "Excel calculations support 15 significant digits. "
                "Use snapshot strings for exact interchange.",
            ],
            [
                "Exception workflow",
                "Frozen snapshot at run; subsequent review does not change comparison or metrics.",
            ],
        ],
        "Metrics": [
            [
                "Metric",
                "Currency",
                "Unit",
                "Numerator",
                "Denominator",
                "Percent",
                "Signed variance",
                "Absolute variance",
            ]
        ],
        "Matches": [
            [
                "Group ID",
                "Keys",
                "Linkage",
                "Comparison",
                "Metric",
                "Currency",
                "Unit",
                "Left",
                "Right",
                "Delta",
                "Absolute variance",
                "Tolerance",
            ]
        ],
        "Exceptions": [["Case ID", "Group ID", "Type", "State at export", "Notes", "Rule ID"]],
        "Data quality": [
            ["Side", "Canonical record", "Invalid", "Quality flags", "Raw source refs"]
        ],
        "Evidence": [
            [
                "Group ID",
                "Side",
                "Canonical",
                "Record ID",
                "Source table",
                "File ID",
                "Sheet",
                "Row",
            ]
        ],
        "Sources": [["Side", "Canonical", "Record ID", "Field", "Value"]],
        "Rules": [
            ["Field", "Value"],
            ["Contract ID", rule.id],
            ["Revision", rule.revision],
            ["Approved by", rule.approved_by],
            ["Approved at", rule.approved_at.isoformat()],
            ["Rule JSON", rule.model_dump_json()],
        ],
    }
    for name in (
        "full_match_rate",
        "linkage_rate",
        "exception_rate",
        "automation_coverage",
        "review_rate",
    ):
        rate = getattr(snapshot, name)
        sections["Metrics"].append(
            [
                name,
                "",
                "CU",
                rate.numerator,
                rate.denominator,
                rate.percent if rate.percent is not None else "N/A",
                "",
                "",
            ]
        )
    for metric in snapshot.monetary:
        sections["Metrics"].append(
            [
                metric.metric_id,
                metric.currency,
                metric.unit,
                "",
                "",
                "",
                metric.signed_variance,
                metric.absolute_variance,
            ]
        )
    for group in all_match_groups(run):
        common = [
            group.id,
            json.dumps(group.key, ensure_ascii=False),
            group.linkage_state,
            group.comparison_state,
        ]
        for comparison in group.comparisons:
            sections["Matches"].append(
                [
                    *common,
                    comparison.metric_id,
                    comparison.currency,
                    comparison.unit,
                    comparison.left_amount,
                    comparison.right_amount,
                    comparison.delta,
                    comparison.absolute_variance,
                    comparison.tolerance,
                ]
            )
        if not group.comparisons:
            sections["Matches"].append([*common, *([None] * 8)])
        for member in group.members:
            for p in member.evidence:
                sections["Evidence"].append(
                    [
                        group.id,
                        member.side,
                        member.canonical_id,
                        member.record_id,
                        p.source_table_id,
                        p.file_id,
                        p.sheet,
                        p.row,
                    ]
                )
    for case in db.scalars(
        select(ExceptionRow).where(ExceptionRow.run_id == run.id).order_by(ExceptionRow.id)
    ):
        view = case_view(case)
        sections["Exceptions"].append(
            [view.id, view.match_group_id, view.type, view.workflow_state, view.notes, view.rule_id]
        )
    actor = Actor("export-worker", run.workspace_id, "editor")
    for side, source in zip(("left", "right"), run.source_refs, strict=True):
        canonical = CanonicalDataset.model_validate(source)
        for record in load_canonical_records(db, actor, canonical):
            sections["Data quality"].append(
                [
                    side,
                    record.id,
                    record.invalid,
                    json.dumps(record.quality_flags),
                    json.dumps([p.model_dump() for p in record.provenance]),
                ]
            )
            for field, value in record.payload.items():
                sections["Sources"].append(
                    [
                        side,
                        canonical.id,
                        record.id,
                        field,
                        value.isoformat() if hasattr(value, "isoformat") else value,
                    ]
                )
    return sections


def write_csv(path: Path, sections: dict[str, list[list[Any]]]) -> None:
    with path.open("x", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        width = max(len(row) for rows in sections.values() for row in rows)
        writer.writerow(["Section", *[f"Column {i}" for i in range(1, width + 1)]])
        for section, rows in sections.items():
            for row in rows:
                writer.writerow(
                    [section, *[safe_text(v) if isinstance(v, str) else v for v in row]]
                )


def write_xlsx(path: Path, sections: dict[str, list[list[Any]]]) -> None:
    workbook = Workbook()
    workbook.remove(workbook.worksheets[0])
    exact: dict[str, dict[str, str]] = {}
    for i, (name, rows) in enumerate(sections.items(), 1):
        sheet = workbook.create_sheet(name)
        exact[f"xl/worksheets/sheet{i}.xml"] = numeric = {}
        for values in rows:
            if any(
                isinstance(v, str)
                and (len(v) > 32767 or re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", v))
                for v in values
            ):
                raise DomainError(
                    "EXPORT_CELL_LIMIT", "Text vượt giới hạn ô XLSX; không cắt hoặc sửa dữ liệu."
                )
            sheet.append(values)
            for cell, value in zip(sheet[sheet.max_row], values, strict=False):
                if isinstance(value, str):
                    cell.data_type = "s"
                    cell.value = safe_text(value)
                elif isinstance(value, Decimal):
                    numeric[cell.coordinate] = str(value)
                    cell.number_format = "#,##0.00############;[Red](#,##0.00############);0.00"
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2563EB")
        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = min(
                55, max(16, max(len(str(c.value or "")) for c in column[:100]) + 2)
            )
    temporary = path.with_suffix(".unpublished.xlsx")
    workbook.save(temporary)
    with ZipFile(temporary) as source, ZipFile(path, "x", compression=ZIP_DEFLATED) as target:
        for name in source.namelist():
            data = source.read(name)
            if name in exact and exact[name]:
                root = ET.fromstring(data)
                for cell in root.findall(f".//{{{NS}}}c"):
                    if cell.attrib.get("r") in exact[name]:
                        cell.set("t", "n")
                        value = cell.find(f"{{{NS}}}v")
                        assert value is not None
                        value.text = exact[name][cell.attrib["r"]]
                data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            target.writestr(name, data)
    temporary.unlink()  # only this known, unpublished local file


def execute_export(job_id: str, token: str, factory: Any, fence: Any, finish: Any) -> None:
    with factory() as db:
        job = fence(db, job_id, token)
        artifact = db.get(ExportRow, job.payload["export_id"])
        assert artifact is not None
        if artifact.status == "SUCCEEDED":
            finish(db, job, {"export_id": artifact.id})
            db.commit()
            return
        run = db.get(RunRow, artifact.run_id)
        assert run is not None
        require_verified(run)
        sections = report_sections(db, run)
        artifact_id, format = artifact.id, artifact.format
        key = f"{run.workspace_id}/exports/{artifact_id}/{identifier()}/report.{format}"
        db.rollback()
    path = storage_path(settings.storage_root, key)
    path.parent.mkdir(parents=True, exist_ok=True)
    if format == "csv":
        write_csv(path, sections)
    elif format == "xlsx":
        write_xlsx(path, sections)
    else:
        raise DomainError("EXPORT_FORMAT_INVALID", "Định dạng báo cáo không được hỗ trợ.")
    digest = checksum(path)
    with factory() as db:
        job = fence(db, job_id, token)
        artifact = db.get(ExportRow, artifact_id)
        assert artifact is not None
        artifact.status, artifact.storage_key, artifact.sha256 = "SUCCEEDED", key, digest
        finish(db, job, {"export_id": artifact_id})
        db.commit()
