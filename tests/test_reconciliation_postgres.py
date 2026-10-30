"""Reuse the existing UUID-schema PostgreSQL harness; no second service or test database."""

import pytest
from alembic import command
from alembic.config import Config
from sqlalchemy import func, select, text
from studio.database import CanonicalRow, RunRow
from test_postgres import approve, drain, pg, project, upload  # noqa: F401, F811

pytestmark = pytest.mark.integration


@pytest.mark.parametrize("pg", ["create_data_foundation"], indirect=True)
def test_additive_migration_preserves_existing_approved_canonical(pg):  # noqa: F811
    client, factory, worker = pg
    project_id = project(client)
    dataset = upload(client, project_id)
    drain(factory, worker)
    table_id = client.get(f"/api/datasets/{dataset['id']}/tables").json()["items"][0]["id"]
    mapping, _ = approve(client, table_id)
    response = client.post(
        f"/api/datasets/{dataset['id']}/normalize",
        json={"mapping_id": mapping["id"]},
        headers={"Idempotency-Key": "migration-canonical"},
    )
    assert response.status_code == 202
    drain(factory, worker)
    before = client.get(f"/api/datasets/{dataset['id']}/artifacts").json()
    assert len(before["canonical_datasets"]) == 1
    command.upgrade(Config("alembic.ini"), "head")
    assert client.get(f"/api/datasets/{dataset['id']}/artifacts").json() == before
    with factory() as db:
        assert db.scalar(select(func.count()).select_from(CanonicalRow)) == 1
        assert db.scalar(select(func.count()).select_from(RunRow)) == 0
        assert db.scalar(text("SELECT version_num FROM alembic_version")) == "create_conversations"
    canonical = before["canonical_datasets"][0]
    preview = client.get(
        f"/api/tables/{table_id}/preview",
        params={"kind": "canonical", "canonical_id": canonical["id"]},
    ).json()
    assert preview["total"] == 4
    assert any(row["record_payload"].get("reference_id") == "000123" for row in preview["items"])


def prepare_pair(pg, *, workbooks=False, grouped=False):  # noqa: F811
    from reconciliation_cases import FIXTURES

    client, factory, worker = pg
    pid = project(client, "Synthetic delivery")
    ids = []
    for filename in ("shipments.csv", "delivery-confirmations.csv"):
        raw = (FIXTURES / filename).read_bytes()
        detail = grouped and filename == "delivery-confirmations.csv"
        grain = "delivery_line" if detail else "delivery"
        if detail:
            raw = raw.replace(
                b"CARRIER-A,000701,2026-10-01,25.00,VND,Parcel A",
                b"CARRIER-A,000701,2026-10-01,10.00,VND,Parcel A\n"
                b"CARRIER-A,000701,2026-10-01,15.00,VND,Parcel A",
            )
        media_type = "text/csv"
        if workbooks:
            import csv
            import io

            from openpyxl import Workbook

            data = list(csv.DictReader(io.StringIO(raw.decode())))
            order = [
                "tracking_id",
                "event_date",
                "delivery_charge",
                "carrier",
                "currency",
                "description",
            ]
            book = Workbook()
            book.active.title = "Notes"
            book.active.append(["Read me"])
            book.active.append(["Synthetic evidence workbook"])
            sheet = book.create_sheet("Deliveries")
            sheet.append(order)
            for item in data:
                sheet.append([item[k] for k in order])
            buffer = io.BytesIO()
            book.save(buffer)
            raw = buffer.getvalue()
            filename += ".xlsx"
            media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        response = client.post(
            "/api/datasets/uploads",
            data={"project_id": pid, "request_ai_mapping": "false"},
            files={"files": (filename, raw, media_type)},
        )
        dataset = response.json()["items"][0]
        drain(factory, worker)
        tables = client.get(f"/api/datasets/{dataset['id']}/tables").json()["items"]
        table = next(t for t in tables if t["sheet"] == "Deliveries") if workbooks else tables[0]
        proposal = client.get(f"/api/tables/{table['id']}/mappings").json()["items"][0]
        fields = [
            ("carrier", "string"),
            ("tracking_id", "string"),
            ("event_date", "date"),
            ("delivery_charge", "decimal"),
            ("currency", "string"),
            ("description", "string"),
        ]
        if workbooks:
            fields = [(f, dict(fields)[f]) for f in order]
        schema = proposal["schema"] | dict(
            entity_kind="delivery",
            grain=grain,
            business_key_fields=["carrier", "tracking_id"],
            fields=[
                dict(
                    id=f,
                    name=f,
                    data_type=t,
                    required=False,
                    **(
                        dict(
                            semantic_metric="delivery_charge",
                            amount_basis="gross",
                            currency="VND",
                            unit="VND",
                        )
                        if t == "decimal"
                        else {}
                    ),
                )
                for f, t in fields
            ],
        )
        decisions = [
            dict(
                source_field=f"c{i}",
                candidate_target=f,
                decision="APPROVED",
                grain_assumption=grain,
                evidence_ids=[f"{table['id']}:c{i}"],
                transformation_candidate=dict(op="parse_decimal", locale="en-US")
                if t == "decimal"
                else dict(op="parse_date", format="%Y-%m-%d", timezone="UTC")
                if t == "date"
                else dict(op="cast_string_preserve"),
            )
            for i, (f, t) in enumerate(fields, 1)
        ]
        response = client.post(
            f"/api/mappings/{proposal['id']}/approve",
            json=dict(schema_definition=schema, decisions=decisions, extraction_reviewed=True),
        )
        assert response.status_code == 200, response.text
        response = client.post(
            f"/api/datasets/{dataset['id']}/normalize",
            json={"mapping_id": response.json()["id"]},
            headers={"Idempotency-Key": filename},
        )
        assert response.status_code == 202, response.text
        drain(factory, worker)
        ids.append(
            client.get(f"/api/datasets/{dataset['id']}/artifacts").json()["canonical_datasets"][0][
                "id"
            ]
        )
    from reconciliation_cases import plan

    p = plan(
        left=ids[0],
        right=ids[1],
        platform="",
        amount="delivery_charge",
        metric="delivery_charge",
        keys=["carrier", "tracking_id"],
        entity="delivery",
        cardinality="ONE_TO_MANY" if grouped else "ONE_TO_ONE",
    )
    return pid, p


def create_approved_run(pg, *, workbooks=False, grouped=False):  # noqa: F811
    client, factory, worker = pg
    pid, p = prepare_pair(pg, workbooks=workbooks, grouped=grouped)
    response = client.post(
        "/api/reconciliation-plans",
        json=dict(project_id=pid, plan=p.model_dump(mode="json")),
        headers={"Idempotency-Key": "delivery-plan:" + pid},
    )
    assert response.status_code == 201, response.text
    plan_id = response.json()["id"]
    rule = client.post(f"/api/reconciliation-contracts/{plan_id}/approve")
    assert rule.status_code == 200, rule.text
    response = client.post(
        "/api/reconciliation-runs",
        json={"contract_id": rule.json()["id"]},
        headers={"Idempotency-Key": "delivery-run"},
    )
    assert response.status_code == 202, response.text
    return pid, p, rule.json(), response.json()


def test_T20_T34_T35_rule_approval_snapshot_exports_and_workflow(pg):  # noqa: F811
    import csv
    import io
    from zipfile import ZipFile

    from openpyxl import load_workbook

    client, factory, worker = pg
    pid, p, rule, accepted = create_approved_run(pg)
    rid = accepted["run_id"]
    assert client.get(f"/api/runs/{rid}/metrics").status_code == 409
    from studio.reconciliation_jobs import execute_reconciliation

    with factory() as db:
        claimed = worker.claim_job(db)
    execute_reconciliation(*claimed, factory, worker._fence, worker._finish)
    run = client.get(f"/api/runs/{rid}").json()
    assert run["status"] == "VERIFIED", run
    preparation = client.get(f"/api/projects/{pid}/summary").json()
    assert preparation["latest_verified_run_id"] == rid
    assert preparation["ready_canonical_count"] == 2
    assert preparation["preparation_action"]["kind"] == "RECONCILE"
    snapshot = client.get(f"/api/runs/{rid}/metrics").json()
    assert (
        client.get(
            "/api/runs", params={"project_id": pid, "source_id": p.left.canonical_id}
        ).json()["total"]
        == 1
    )
    assert (
        client.get("/api/runs", params={"project_id": pid, "period_start": "2026-11-01"}).json()[
            "total"
        ]
        == 0
    )
    assert snapshot["monetary"][0]["absolute_variance"] == "2.75"
    assert snapshot["full_match_rate"] == dict(numerator=1, denominator=3, percent="33.33")
    duplicate = client.post(
        "/api/reconciliation-runs",
        json={"contract_id": rule["id"]},
        headers={"Idempotency-Key": "another"},
    ).json()
    assert duplicate["run_id"] == rid
    cases = client.get(f"/api/runs/{rid}/exceptions").json()
    case = next(c for c in cases["items"] if c["type"] == "AMOUNT_MISMATCH")
    assert (
        client.patch(
            f"/api/exceptions/{case['id']}/review",
            json=dict(workflow_state="RESOLVED", notes="Reviewed synthetic delta"),
        ).status_code
        == 200
    )
    assert client.get(f"/api/runs/{rid}/metrics").json() == snapshot
    evidence = client.get(f"/api/runs/{rid}/groups/{case['match_group_id']}/evidence")
    assert evidence.status_code == 200, evidence.text
    assert sum(s["member_count"] for s in evidence.json()["sources"]) == 2
    for format in ("csv", "xlsx"):
        exported = client.post(
            f"/api/runs/{rid}/exports?format={format}", headers={"Idempotency-Key": format}
        )
        assert exported.status_code == 202, exported.text
        drain(factory, worker)
        info = client.get(f"/api/exports/{exported.json()['export_id']}").json()
        assert info["status"] == "SUCCEEDED", info
        downloaded = client.get(info["download_url"])
        assert downloaded.status_code == 200
        if format == "csv":
            rows = list(csv.reader(io.StringIO(downloaded.content.decode("utf-8-sig"))))
            stored = next(r[2] for r in rows if r[:2] == ["Summary", "Snapshot JSON"])
        else:
            book = load_workbook(io.BytesIO(downloaded.content), data_only=False)
            stored = next(r[1].value for r in book["Summary"] if r[0].value == "Snapshot JSON")
            assert all(c.data_type != "f" for sheet in book for row in sheet for c in row)
            with ZipFile(io.BytesIO(downloaded.content)) as archive:
                assert any(
                    b"2.75" in archive.read(n)
                    for n in archive.namelist()
                    if n.startswith("xl/worksheets/")
                )
        import json

        assert json.loads(stored) == snapshot
    assert (
        client.post(
            "/api/reconciliation-runs",
            json={"contract_id": rule["plan"]["left"]["canonical_id"]},
            headers={"Idempotency-Key": "unapproved"},
        ).status_code
        == 404
    )


def test_T33_crash_before_publish_and_failed_verification_no_metrics(pg, monkeypatch):  # noqa: F811
    from studio import reconciliation_jobs

    client, factory, worker = pg
    _, _, _, accepted = create_approved_run(pg)
    original = reconciliation_jobs.write_match_artifact

    def fail(*args, **kwargs):
        original(*args, **kwargs)
        raise RuntimeError("synthetic crash")

    monkeypatch.setattr(reconciliation_jobs, "write_match_artifact", fail)
    with factory() as db:
        claim = worker.claim_job(db)
    worker.process_job(*claim)
    assert client.get(f"/api/runs/{accepted['run_id']}/metrics").status_code == 409
    with factory() as db:
        from studio.database import JobRow, now

        job = db.get(JobRow, accepted["job_id"])
        job.available_at = now()
        db.commit()
    monkeypatch.setattr(reconciliation_jobs, "write_match_artifact", original)
    drain(factory, worker)
    assert client.get(f"/api/runs/{accepted['run_id']}/metrics").status_code == 200


def test_failed_verification_blocks_metrics_exports_and_cross_run_evidence(pg, monkeypatch):  # noqa: F811
    from decimal import Decimal

    from studio import reconciliation_jobs

    client, factory, worker = pg
    _, _, _, accepted = create_approved_run(pg)
    original = reconciliation_jobs.match_records

    def corrupt(*args):
        result = original(*args)
        next(g for g in result.groups if g.comparisons).comparisons[0].left_amount = Decimal(
            "99999"
        )
        return result

    monkeypatch.setattr(reconciliation_jobs, "match_records", corrupt)
    drain(factory, worker)
    rid = accepted["run_id"]
    run = client.get(f"/api/runs/{rid}").json()
    assert run["status"] == "VERIFICATION_FAILED" and run["metric_snapshot"] is None
    assert any(not c["passed"] for c in run["checks"])
    assert client.get(f"/api/runs/{rid}/metrics").status_code == 409
    assert (
        client.post(f"/api/runs/{rid}/exports", headers={"Idempotency-Key": "bad"}).status_code
        == 409
    )


def test_reconciliation_permissions_scope_and_template_drift(pg):  # noqa: F811
    from studio.database import CanonicalRow, MembershipRow

    client, factory, worker = pg
    pid, p, rule, accepted = create_approved_run(pg)
    drain(factory, worker)
    rid = accepted["run_id"]
    assert client.get(f"/api/runs/{rid}/groups/not-a-group/evidence").status_code == 404
    other = project(client, "other project")
    response = client.post(
        "/api/reconciliation-plans",
        json=dict(project_id=other, plan=p.model_dump(mode="json")),
        headers={"Idempotency-Key": "foreign"},
    )
    assert response.status_code == 422
    payload = dict(
        left_canonical_id=p.left.canonical_id,
        right_canonical_id=p.right.canonical_id,
        period_start="2026-11-01",
        period_end="2026-12-01",
        left_coverage_start="2026-11-01",
        left_coverage_end="2026-12-01",
        right_coverage_start="2026-11-01",
        right_coverage_end="2026-12-01",
        completeness="CONFIRMED",
        snapshot_at="2026-12-02T00:00:00Z",
    )
    response = client.post(f"/api/reconciliation-contracts/{rule['id']}/reuse", json=payload)
    assert response.status_code == 201, response.text
    assert (
        response.json()["contract_id"] is None
        and response.json()["proposal"]["status"] == "PROPOSED"
    )
    with factory() as db:
        membership = db.get(MembershipRow, ("demo", "test-owner"))
        membership.role = "viewer"
        db.commit()
    response = client.post(
        "/api/reconciliation-runs",
        json={"contract_id": rule["id"]},
        headers={"Idempotency-Key": "viewer"},
    )
    assert response.status_code == 403
    with factory() as db:
        membership = db.get(MembershipRow, ("demo", "test-owner"))
        membership.role = "owner"
        source = db.get(CanonicalRow, p.left.canonical_id)
        source.stale = True
        db.commit()
    assert (
        client.post(
            "/api/reconciliation-runs",
            json={"contract_id": rule["id"]},
            headers={"Idempotency-Key": "stale"},
        ).status_code
        == 422
    )


def test_planner_accepts_sources_without_keys_but_cannot_execute(pg):  # noqa: F811
    from studio.database import PlanRow

    client, factory, worker = pg
    pid, p = prepare_pair(pg)
    response = client.post(
        "/api/reconciliation-plans/ai",
        json=dict(
            project_id=pid,
            request=dict(
                left={"canonical_id": p.left.canonical_id},
                right={"canonical_id": p.right.canonical_id},
                intent="Suggest keys between these sources",
            ),
        ),
        headers={"Idempotency-Key": "no-keys"},
    )
    assert response.status_code == 202, response.text
    # No worker/provider execution: a queued proposal grants no execution authority.
    assert (
        client.post(
            "/api/reconciliation-runs",
            json={"contract_id": response.json()["job_id"]},
            headers={"Idempotency-Key": "not-approved"},
        ).status_code
        == 404
    )
    with factory() as db:
        assert db.scalar(select(func.count()).select_from(RunRow)) == 0
        row = PlanRow(
            workspace_id="demo",
            project_id=pid,
            origin="GEMINI",
            created_by="test-owner",
            proposal={
                "status": "NEEDS_REVIEW",
                "plan": None,
                "missing_requirements": ["Confirm coverage"],
            },
        )
        db.add(row)
        db.commit()
        plan_id = row.id
    denied = client.post(f"/api/reconciliation-contracts/{plan_id}/approve")
    assert denied.status_code == 422 and denied.json()["code"] == "PLAN_NEEDS_REVIEW"


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_export_first_generation_is_immutable_after_review(pg, format):  # noqa: F811
    import csv
    import io

    from openpyxl import load_workbook

    client, factory, worker = pg
    _, _, _, accepted = create_approved_run(pg)
    drain(factory, worker)
    rid = accepted["run_id"]
    snapshot = client.get(f"/api/runs/{rid}/metrics").json()
    endpoint = f"/api/runs/{rid}/exports?format={format}"
    created = client.post(endpoint, headers={"Idempotency-Key": "first"}).json()
    drain(factory, worker)
    info = client.get(f"/api/exports/{created['export_id']}").json()
    before = client.get(info["download_url"]).content
    assert info["review_state_scope"] == "FIRST_GENERATION" and info["review_state_as_of"]
    if format == "csv":
        rows = list(csv.reader(io.StringIO(before.decode("utf-8-sig"))))
        summary = {r[1]: r[2] for r in rows if r[0] == "Summary"}
        states = [r[4] for r in rows if r[0] == "Exceptions"][1:]
    else:
        book = load_workbook(io.BytesIO(before))
        summary = dict(book["Summary"].values)
        states = [r[3] for r in list(book["Exceptions"].values)[1:]]
    assert summary["Report generated at"] == info["review_state_as_of"]
    assert "FIRST_GENERATION" in summary["Review state scope"]
    assert set(states) == {"OPEN"}
    case = client.get(f"/api/runs/{rid}/exceptions").json()["items"][0]
    assert (
        client.patch(
            f"/api/exceptions/{case['id']}/review",
            json={"workflow_state": "RESOLVED", "notes": "Reviewed after immutable export"},
        ).status_code
        == 200
    )
    second = client.post(endpoint, headers={"Idempotency-Key": "again"}).json()
    assert second["export_id"] == created["export_id"]
    assert client.get(info["download_url"]).content == before
    assert client.get(f"/api/runs/{rid}/metrics").json() == snapshot


def test_readiness_duplicates_use_engine_period_partition(pg, monkeypatch):  # noqa: F811
    from datetime import UTC, datetime

    from reconciliation_cases import records
    from studio import reconciliation_api

    client, _, _ = pg
    pid, p = prepare_pair(pg)
    rows = records("shipments.csv")
    earlier = rows[0].model_copy(deep=True)
    earlier.id = "earlier"
    earlier.payload["event_date"] = datetime(2026, 9, 20, tzinfo=UTC)
    monkeypatch.setattr(
        reconciliation_api, "load_canonical_records", lambda *args: [*rows, earlier]
    )
    response = client.post(
        "/api/reconciliation-plans/validate",
        json=dict(project_id=pid, plan=p.model_dump(mode="json")),
    )
    assert response.status_code == 200, response.text
    assert response.json()["sources"][0]["duplicate_key_rows"] == 0
    assert response.json()["sources"][0]["out_of_scope_records"] == 1
    duplicate = rows[0].model_copy(deep=True)
    duplicate.id = "duplicate-in-period"
    missing_date = rows[0].model_copy(deep=True)
    missing_date.id, missing_date.payload["event_date"] = "missing-date", None
    invalid_key = rows[0].model_copy(deep=True)
    invalid_key.id, invalid_key.payload["tracking_id"] = "invalid-key", ""
    outside = rows[0].model_copy(deep=True)
    outside.id, outside.payload["carrier"] = "other-carrier", "OTHER"
    all_rows = [*rows, earlier, duplicate, missing_date, invalid_key, outside]
    monkeypatch.setattr(reconciliation_api, "load_canonical_records", lambda *args: all_rows)
    for slot in (p.left, p.right):
        slot.filters, slot.completeness = {"carrier": "CARRIER-A"}, "INCOMPLETE"
    p.allow_partial = True
    response = client.post(
        "/api/reconciliation-plans/validate",
        json=dict(project_id=pid, plan=p.model_dump(mode="json")),
    )
    source = response.json()["sources"][0]
    assert source["duplicate_key_rows"] == 1 and source["in_scope_records"] == 4
    assert source["invalid_records"] == 2 and source["out_of_scope_records"] == 2
    assert source["coverage_complete"] is False and source["duplicate_population"] == "IN_SCOPE"


def test_planner_exhausted_durable_budget_no_cloud_manual_fallback(pg, monkeypatch):  # noqa: F811
    import httpx
    from pydantic import SecretStr
    from studio import reconciliation_jobs
    from studio.ai_gateway import GeminiProvider
    from studio.config import settings
    from studio.database import AuditRow, JobRow, WorkspaceRow

    client, factory, worker = pg
    pid, p = prepare_pair(pg)
    monkeypatch.setattr(settings, "allow_external_ai", True)
    monkeypatch.setattr(settings, "external_ai_request_budget", 10)
    monkeypatch.setattr(settings, "gemini_api_key", SecretStr("synthetic-test-key"))
    with factory() as db:
        workspace = db.get(WorkspaceRow, "demo")
        workspace.policy = workspace.policy | dict(
            mode="CLOUD_ALLOWED", allow_external_ai=True, external_request_budget=10
        )
        for i in range(10):
            db.add(
                AuditRow(
                    workspace_id="demo",
                    actor_id="test-owner",
                    action="GEMINI_REQUEST_RESERVED",
                    resource_id=str(i),
                )
            )
        db.commit()
    calls = []

    def handle(request):
        calls.append(request)
        pytest.fail("Exhausted budget must never call a provider")

    provider = GeminiProvider(settings, httpx.Client(transport=httpx.MockTransport(handle)))
    monkeypatch.setattr(reconciliation_jobs, "GeminiProvider", lambda config: provider)
    response = client.post(
        "/api/reconciliation-plans/ai",
        json=dict(
            project_id=pid,
            request=dict(
                left={"canonical_id": p.left.canonical_id},
                right={"canonical_id": p.right.canonical_id},
            ),
        ),
        headers={"Idempotency-Key": "exhausted-planning"},
    )
    assert response.status_code == 202, response.text
    drain(factory, worker)
    with factory() as db:
        job = db.get(JobRow, response.json()["job_id"])
        assert job.status == "BLOCKED" and job.error_code == "CLOUD_BUDGET_EXHAUSTED"
        assert (
            db.scalar(
                select(func.count())
                .select_from(AuditRow)
                .where(AuditRow.action == "GEMINI_REQUEST_RESERVED")
            )
            == 10
        )
    assert calls == []
    manual = client.post(
        "/api/reconciliation-plans",
        json=dict(project_id=pid, plan=p.model_dump(mode="json")),
        headers={"Idempotency-Key": "manual-after-budget"},
    )
    assert manual.status_code == 201, manual.text


def test_source_evidence_real_columns_multisheet_prefix_and_viewer(pg):  # noqa: F811
    from studio.database import MembershipRow

    client, factory, worker = pg
    _, _, _, accepted = create_approved_run(pg, workbooks=True, grouped=True)
    drain(factory, worker)
    rid = accepted["run_id"]
    group = client.get(f"/api/runs/{rid}/matches?comparison=AMOUNT_MISMATCH").json()["items"][0]
    _, _, _, foreign = create_approved_run(pg)
    drain(factory, worker)
    foreign_group = client.get(f"/api/runs/{foreign['run_id']}/matches").json()["items"][0]
    assert client.get(f"/api/runs/{rid}/groups/{foreign_group['id']}/evidence").status_code == 404
    with factory() as db:
        db.get(MembershipRow, ("demo", "test-owner")).role = "viewer"
        db.commit()
    evidence = client.get(f"/api/runs/{rid}/groups/{group['id']}/evidence").json()
    for source, amount in zip(evidence["sources"], ("12.75", "10.00"), strict=True):
        record = source["records"][0]
        refs = record["field_evidence"]
        assert refs["tracking_id"][0]["column"] == "c1"
        ref = refs["delivery_charge"][0]
        assert ref["column"] == "c3" and ref["sheet"] == "Deliveries"
        raw = client.get(f"/api/evidence/{ref['id']}", params={"row": ref["row"]}).json()
        assert raw["sheet"] == "Deliveries" and raw["items"][0]["c3"] == amount
        id_ref = refs["tracking_id"][0]
        raw_id = client.get(f"/api/evidence/{id_ref['id']}", params={"row": id_ref["row"]}).json()
        assert raw_id["items"][0]["c1"] == "000702"
    linked = client.get(f"/api/runs/{rid}/matches?comparison=EQUAL").json()["items"][0]
    grouped_evidence = client.get(f"/api/runs/{rid}/groups/{linked['id']}/evidence").json()
    assert [s["member_count"] for s in grouped_evidence["sources"]] == [1, 2]
    refs = [
        r["field_evidence"]["delivery_charge"][0] for r in grouped_evidence["sources"][1]["records"]
    ]
    assert len({ref["row"] for ref in refs}) == 2
    assert {
        client.get(f"/api/evidence/{ref['id']}", params={"row": ref["row"]}).json()["items"][0][
            "c3"
        ]
        for ref in refs
    } == {"10.00", "15.00"}
    assert (
        client.patch(
            "/api/exceptions/nonexistent/review",
            json={"workflow_state": "RESOLVED", "notes": "viewer denied"},
        ).status_code
        == 403
    )
