"""Create tiny, original synthetic document fixtures using bundled artifact Python.

No provider calls. Each document is one page with a manually specified row oracle.
Run separately with --format pdf or --format docx after the artifact-operation marker.
"""

import argparse
from pathlib import Path


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--format", choices=["pdf", "docx"], required=True)
    args = parser.parse_args()
    root = Path(__file__).resolve().parents[1] / "runtime" / "document-acceptance"
    root.mkdir(parents=True, exist_ok=True)
    target = root / f"synthetic-transactions.{args.format}"
    if target.exists():
        raise SystemExit("Fixture already exists; preserve the original bytes and provider reuse.")
    headers = ["Transaction ID", "Date", "Amount", "Description"]
    rows = (
        [
            ["000123", "2026-10-01", "1234.50", "Paid"],
            ["000124", "2026-10-02", "-20.00", "Refund"],
            ["000125", "2026-10-03", "", "Pending"],
        ]
        if args.format == "pdf"
        else [
            ["000701", "2026-10-04", "1.234,50", "Priority A"],
            ["000702", "2026-10-05", "0,00", "Standard"],
            ["000703", "2026-10-06", "-12,75", "Refund"],
        ]
    )
    title = "Transaction register"
    intro = "Transactions for October 2026. Amounts use the printed number format."
    note = (
        "Note: the pending transaction has no amount."
        if args.format == "pdf"
        else ("Note: amounts use a comma decimal separator and a dot thousands separator.")
    )
    if args.format == "pdf":
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

        document = SimpleDocTemplate(
            str(target),
            pagesize=letter,
            title=title,
            author="AI Reconciliation Studio",
            leftMargin=54,
            rightMargin=54,
            topMargin=54,
            bottomMargin=54,
        )
        styles = getSampleStyleSheet()
        table = Table([headers, *rows], colWidths=[110, 105, 95, 140], repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D9D9D9")),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E2E8F0")),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, -1), 11),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 8),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                    ("TOPPADDING", (0, 0), (-1, -1), 10),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
                ]
            )
        )
        document.build(
            [
                Paragraph(title, styles["Title"]),
                Spacer(1, 16),
                Paragraph(intro, styles["Normal"]),
                Spacer(1, 20),
                table,
                Spacer(1, 20),
                Paragraph(note, styles["Normal"]),
            ]
        )
    else:
        from docx import Document
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt, RGBColor

        document = Document()
        section = document.sections[0]
        section.page_width, section.page_height = Inches(8.5), Inches(11)
        section.left_margin = section.right_margin = Inches(0.75)
        section.top_margin = section.bottom_margin = Inches(0.75)
        document.core_properties.author = "AI Reconciliation Studio"
        for name in ["Normal", "Title"]:
            style = document.styles[name]
            style.font.name = "Arial"
            style.font.color.rgb = RGBColor(0, 0, 0)
        document.styles["Normal"].font.size = Pt(11)
        document.add_paragraph(title, "Title")
        document.add_paragraph(intro)
        table = document.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        table.autofit = False
        for column, width in zip(table.columns, [1.6, 1.5, 1.4, 2.3], strict=True):
            column.width = Inches(width)
        for i, values in enumerate([headers, *rows]):
            cells = table.rows[0].cells if i == 0 else table.add_row().cells
            for cell, value in zip(cells, values, strict=True):
                cell.text = value
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcpr = cell._tc.get_or_add_tcPr()
                borders = OxmlElement("w:tcBorders")
                for side in ["top", "left", "bottom", "right"]:
                    edge = OxmlElement(f"w:{side}")
                    for key, val in {"val": "single", "sz": "4", "color": "D9D9D9"}.items():
                        edge.set(qn(f"w:{key}"), val)
                    borders.append(edge)
                tcpr.append(borders)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_before = Pt(8)
                    paragraph.paragraph_format.space_after = Pt(8)
                if i == 0:
                    shading = OxmlElement("w:shd")
                    shading.set(qn("w:fill"), "E2E8F0")
                    tcpr.append(shading)
                    cell.paragraphs[0].runs[0].bold = True
        document.add_paragraph(note)
        document.save(target)
    print("Created synthetic fixture:", target.name)


if __name__ == "__main__":
    main()
